import win32com.client as win32
import openpyxl
import os
import shutil # <-- needed to copy the template file
import pywintypes # <-- needed to convert dates for Excel COM
from datetime import datetime, date, time, timedelta

# Data file paths
USER_HOME = os.path.expanduser("~")

# file_0_path = USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN\280_HR\004_Training\100_Training\Training_Evidence_Scan_Test.xlsx"
file_0_path = r"C:\Training HR scripts\Output\Training_Evidence_Scan_Test.xlsx"
file_1_path = USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN\280_HR\004_Training\100_Training\Training_Record_Template.xlsx"
# file_2_path = USER_HOME + r"\RRC power solutions\RRC VN - Documents\280_HR\104_ Employee_Management_Table\02. Employee management\RRC_Employee_Master_Data_2025-07-10.xlsx"
file_2_path = r"C:\Training HR scripts\RRC_Employee_Master_Data_2025-07-10.xlsx"
file_3_path = USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN\280_HR\004_Training\100_Training\Training_Matrix.xlsx"

# Log file path in same folder like file1 = template
log_path = os.path.join(os.path.dirname(file_1_path), "Training_Record.log")

# Output file with current date & time
timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
output_filename = f"Training_Record_{timestamp}.xlsx"
# output_path = os.path.join(os.path.dirname(file_1_path), output_filename)
output_dir = r"C:\Training HR scripts\Output"
# output_dir = USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN\280_HR\004_Training\100_Training"
output_path = os.path.join(output_dir, output_filename)

# Make a copy of the template to work on, so original is not modified
shutil.copyfile(file_1_path, output_path)

# *** START LOGGING ***
start_time = datetime.now()
log_lines = []
log_lines.append("*** Training Record Processing Started ***")
log_lines.append(f"Start Time: {start_time.strftime('%Y-%m-%d %H:%M:%S')}")

def log(message):
    print(message)
    log_lines.append(message)

# Problem with date and time zone: every date was 1 day earlier due to VN time zone: now ensure datetime is set to midnight without time zone offset
def normalize_date(cell_value):
    if isinstance(cell_value, datetime):
        return datetime.combine(cell_value.date(), time(0, 0))
    return cell_value 

# *** STEP 0: Load training evidence data from File_0 ***
log("*** STEP 0: Loading training evidence data...")
wb0 = openpyxl.load_workbook(file_0_path, data_only=True)
ws0 = wb0.active

training_data = {}  # Dictionary to store training data by employee ID
row = 2  # Data starts in line 2

# Read data row by row starting from row 2 until no data found
while True:
    employee_id = ws0[f"L{row}"].value  # Employee_ID is in column L
    if employee_id is None:
        break
    
    # Initialize employee training list if not exists
    if str(employee_id) not in training_data:
        training_data[str(employee_id)] = []
    
    # Extract training data for this employee
    filename = ws0[f"D{row}"].value[5:]
    training_record = {
        "Training_No": ws0[f"E{row}"].value,  # Column E
        "Expiration_Date": normalize_date(ws0[f"H{row}"].value),  # Column H
        "Training_Evidence_No": ws0[f"B{row}"].value,  # Column B (assuming this is the correct column)
        "Hyperlink": f'=HYPERLINK("https://rrcpowersolutions.sharepoint.com/sites/RRCVN/Shared Documents/999_SHARE_VN/280_HR/004_Training/100_Training/Training_Evidence_Scan/PDF_to_import/{filename}", "Open {filename}")',  # Column D  
        "Training_Compliance_Status": ws0[f"I{row}"].value,  # Column I (you may need to adjust this)
        "Manual_Approval_Status": ws0[f"J{row}"].value,  # Column J
    }
    
    training_data[str(employee_id)].append(training_record)
    row += 1

log(f"Loaded training data for {len(training_data)} employees from Training_Evidence_Scan file.")
log("Training data preview:")
for emp_id, records in training_data.items():
    log(f" Employee {emp_id}: {len(records)} training records")

# *** STEP 1: Load employee master data from file 2 ***
log("*** STEP 1: Loading employee master data...")
wb2 = openpyxl.load_workbook(file_2_path, data_only=True)
ws2 = wb2.active

employee_data = {}
row = 5

# Read data row by row starting from row 5 until no Employee_ID found
while True:
    emp_id = ws2[f"A{row}"].value
    if emp_id is None:
        break
    employee_data[str(emp_id)] = {
        "Full Name": ws2[f"B{row}"].value,
        "Department": ws2[f"C{row}"].value,
        "Position": ws2[f"D{row}"].value,
        "Training Group": ws2[f"E{row}"].value,
        "On-Board-Date": normalize_date(ws2[f"F{row}"].value),
        "Exit-Date": normalize_date(ws2[f"AZ{row}"].value),
        "In_Maternity_Leave": ws2[f"BF{row}"].value, # <-- NEW
    }
    row += 1

log(f"Loaded {len(employee_data)} employees from master data.")
log("Employee data preview (from Employee_Master_Data file):")
for emp_id, info in employee_data.items():
    log(f" ID: {emp_id}, Name: {info['Full Name']}, Dept: {info['Department']}, Pos: {info['Position']}, Group: {info['Training Group']}, On-Board: {info['On-Board-Date']}, Exit: {info['Exit-Date']}, In_Maternity_Leave: {info['In_Maternity_Leave']}")

# *** STEP 2: Open the Excel template workbook via Win32COM ***
log("*** STEP 2: Start Excel ...")
excel = win32.gencache.EnsureDispatch("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
wb = excel.Workbooks.Open(output_path)

# *** STEP 3: Copy TEMPLATE sheet for each employee and add employee data ***
template = wb.Sheets("TEMPLATE")
log("*** STEP 3: Create a sheet for every RRC employee ...")

def calculate_training_status(compliance_status, manual_status):
    """Calculate training status based on compliance and manual approval status"""
    if compliance_status == "NOK" and manual_status == "OK":
        return "OK"
    elif compliance_status == "Recertified":
        return "OK"
    elif compliance_status == "OK":
        return "OK"
    else:
        return compliance_status  # Return original status if none of the conditions match

for emp_id, info in employee_data.items():
    log(f" Create sheet for: {emp_id}")
    template.Copy(After=wb.Sheets(wb.Sheets.Count)) # Duplicate TEMPLATE
    new_sheet = wb.Sheets(wb.Sheets.Count) # Get reference to the new sheet
    new_sheet.Name = str(emp_id)
    
    # Fill in personal data in B4 to H4
    new_sheet.Range("B4").Value = info["Full Name"]
    new_sheet.Range("C4").Value = info["Training Group"]
    new_sheet.Range("D4").Value = info["Department"]
    new_sheet.Range("E4").Value = info["Position"]
    
    if isinstance(info["On-Board-Date"], date):
        new_sheet.Range("F4").Value = info["On-Board-Date"].strftime("%Y-%m-%d")
        new_sheet.Range("F4").NumberFormat = "yyyy-mm-dd"
    else:
        new_sheet.Range("F4").Value = info["On-Board-Date"]
    
    if isinstance(info["Exit-Date"], date):
        new_sheet.Range("G4").Value = info["Exit-Date"].strftime("%Y-%m-%d")
        new_sheet.Range("G4").NumberFormat = "Short Date"
    else:
        new_sheet.Range("G4").Value = info["Exit-Date"]
    
    # In_Maternity_Leave to H4
    new_sheet.Range("H4").Value = info["In_Maternity_Leave"]
    
    # *** Add training data from STEP 0 ***
    if str(emp_id) in training_data:
        employee_training_records = training_data[str(emp_id)]
        log(f"  Adding {len(employee_training_records)} training records for employee {emp_id}")
        
        # Prepare all data in a single 2D array for bulk insertion
        all_training_data = []
        current_date = datetime.now().date()
        
        for i, record in enumerate(employee_training_records):
            # Prepare row data [A, C, D, E, F, G, H, I] - Skip column B (has permanent formula)
            row_data = []
            
            # a.) Training_No (Column A)
            row_data.append(record["Training_No"])
            
            # c.) Expiration date (Column C)
            if isinstance(record["Expiration_Date"], date):
                row_data.append(record["Expiration_Date"].strftime("%Y-%m-%d"))
            else:
                row_data.append(record["Expiration_Date"])
            
            # d.) Days to expiration calculation (Column D)
            if isinstance(record["Expiration_Date"], (date, datetime)):
                if isinstance(record["Expiration_Date"], datetime):
                    exp_date = record["Expiration_Date"].date()
                else:
                    exp_date = record["Expiration_Date"]
                days_diff = (exp_date - current_date).days
                row_data.append(days_diff)
            else:
                row_data.append("")
            
            # e.) Training Evidence No (Column E)
            row_data.append(record["Training_Evidence_No"])
            
            # f.) Hyperlink (Column F)
            row_data.append(record["Hyperlink"])
            
            # g.) Training Compliance Status (Column G)
            row_data.append(record["Training_Compliance_Status"])
            
            # h.) Manual Approval Status (Column H)
            row_data.append(record["Manual_Approval_Status"])
            
            # i.) Calculated Training Status (Column I)
            training_status = calculate_training_status(
                record["Training_Compliance_Status"], 
                record["Manual_Approval_Status"]
            )
            row_data.append(training_status)
            
            all_training_data.append(row_data)
        
        # Write data in 1 operations (instead of excel operation x-times) - skip column B (has permanent formulas)
        if all_training_data:
            start_row = 8
            end_row = start_row + len(all_training_data) - 1
            
            try:
                # Prepare separate arrays for each column (excluding B)
                col_a_data = [[row[0]] for row in all_training_data]  # Training_No
                col_c_data = [[row[1]] for row in all_training_data]  # Expiration date
                col_d_data = [[row[2]] for row in all_training_data]  # Days to expiration
                col_e_data = [[row[3]] for row in all_training_data]  # Evidence No
                col_f_data = [[row[4]] for row in all_training_data]  # Hyperlink
                col_g_data = [[row[5]] for row in all_training_data]  # Compliance Status
                col_h_data = [[row[6]] for row in all_training_data]  # Manual Status
                col_i_data = [[row[7]] for row in all_training_data]  # Training Status
                
                # Bulk write to each column (skip B)
                new_sheet.Range(f"A{start_row}:A{end_row}").Value = col_a_data
                new_sheet.Range(f"C{start_row}:C{end_row}").Value = col_c_data
                new_sheet.Range(f"D{start_row}:D{end_row}").Value = col_d_data
                new_sheet.Range(f"E{start_row}:E{end_row}").Value = col_e_data
                new_sheet.Range(f"F{start_row}:F{end_row}").Value = col_f_data
                new_sheet.Range(f"G{start_row}:G{end_row}").Value = col_g_data
                new_sheet.Range(f"H{start_row}:H{end_row}").Value = col_h_data
                new_sheet.Range(f"I{start_row}:I{end_row}").Value = col_i_data
                
                log(f"  Successfully added {len(all_training_data)} training records to sheet {emp_id}")
            except Exception as e:
                log(f"  Error with bulk operation for {emp_id}: {e}")
                log(f"  Falling back to individual cell operations...")
                
                # Fallback: write data row by row if bulk fails
                for i, row_data in enumerate(all_training_data):
                    current_row = start_row + i
                    try:
                        # Write to individual columns (skip B)
                        new_sheet.Range(f"A{current_row}").Value = row_data[0]  # Training_No
                        new_sheet.Range(f"C{current_row}").Value = row_data[1]  # Expiration date
                        new_sheet.Range(f"D{current_row}").Value = row_data[2]  # Days to expiration
                        new_sheet.Range(f"E{current_row}").Value = row_data[3]  # Evidence No
                        new_sheet.Range(f"F{current_row}").Value = row_data[4]  # Hyperlink
                        new_sheet.Range(f"G{current_row}").Value = row_data[5]  # Compliance Status
                        new_sheet.Range(f"H{current_row}").Value = row_data[6]  # Manual Status
                        new_sheet.Range(f"I{current_row}").Value = row_data[7]  # Training Status
                    except Exception as row_error:
                        log(f"    Error setting row {current_row}: {row_error}")
                
                log(f"  Completed fallback operations for sheet {emp_id}")
    else:
        log(f"  No training data found for employee {emp_id}")

# *** STEP 4: Update the OVERVIEW sheet ***
overview = wb.Sheets("OVERVIEW")
log("*** STEP 4: Updating OVERVIEW sheet...")

# Clear previous entries (clean template makes this optional, but we keep it safe)
overview.Range("A3:J10000").ClearContents() # <-- widened to J

# Define overview data with row number in column A
overview_data = [] #list to write all data in which is needed for the overview sheet
for idx, (emp_id, info) in enumerate(employee_data.items(), start=1):
    hyperlink_formula = f'=HYPERLINK("#\'{emp_id}\'!A1", "Go to sheet: {emp_id}")'
    overview_data.append([ #new line is created and add with the data A...J
        idx, # A: Consecutive numbering
        emp_id, # B: Employee_ID
        hyperlink_formula, # C: Hyperlink to jump to the related employee sheet
        info["Full Name"], # D
        info["Training Group"], # E
        info["Department"], # F
        info["Position"], # G
        info["On-Board-Date"].strftime("%Y-%m-%d") if isinstance(info["On-Board-Date"], date) else info["On-Board-Date"], # H
        info["Exit-Date"].strftime("%Y-%m-%d") if isinstance(info["Exit-Date"], date) else info["Exit-Date"], # I
        info["In_Maternity_Leave"], # J <-- NEW
    ])

# Write the full table to Excel
end_row = 2 + len(overview_data)
log("Data prepared for OVERVIEW sheet:")
for row in overview_data:
    log(f" {row}")
overview.Range(f"A3:J{end_row}").Value = overview_data # <-- widened to J
log(f"Overview sheet updated with {len(employee_data)} entries.")

# *** STEP 4.5: Backfill required trainings from Training_Matrix for each active employee ***

# Load Training_Matrix
try:
    wb3 = openpyxl.load_workbook(file_3_path, data_only=True)
    try:
        ws3 = wb3["Training Matrix"]
    except KeyError:
        ws3 = wb3.active
    log(f'Loaded Training_Matrix from: {file_3_path} (sheet="{ws3.title}")')
except Exception as e:
    ws3 = None
    log(f"WARNING: Could not open Training_Matrix: {e}")

def _to_str(v):
    """
    Canonicalize Excel/OOXML values to string:
    - If float with .0, drop the .0 (e.g., 230019.0 -> "230019", 46.0 -> "46")
    - If float with decimals, trim trailing zeros and trailing dot
    - Strip whitespace; return None for empty strings
    """
    if v is None:
        return None
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        s = str(v).rstrip('0').rstrip('.')
        return s
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s if s != "" else None

# --- Active-Flag map (Active Flag:  training active=1 / not active=0) from Training_Matrix (Row 4 E4..; Training No. headers in Row 6 E6..) ---
ACTIVE_FLAG_BY_TRAINING = {}
if ws3 is not None:
    max_col = ws3.max_column
    for c in range(5, max_col + 1):  # E = 5 .. end
        tn = _to_str(ws3.cell(row=6, column=c).value)   # Training_No (row 6)
        af = _to_str(ws3.cell(row=4, column=c).value)   # Active Flag (row 4)
        if tn:
            # Only "1" is active; anything else (0/2/blank) is treated as inactive
            ACTIVE_FLAG_BY_TRAINING[tn] = "1" if af == "1" else "0"




def get_required_training_numbers(group_val):
    """
    Find the row for the given Training_Group in column D (starting at row 7),
    then scan across row 6 for Training_No. headers (E6 -> ...).
    For the group's row, if cell != '0' (and not blank), include that Training_No.
    """
    if ws3 is None:
        return []

    target = _to_str(group_val)
    if target is None:
        return []

    # Find the group's row in column D, starting at row 7
    row_index = None
    for r in range(7, ws3.max_row + 1):
        val = _to_str(ws3.cell(row=r, column=4).value)  # column D
        if val == target:
            row_index = r
            break

    if row_index is None:
        log(f"  WARNING: Training group '{target}' not found in Training_Matrix.")
        return []

    required = []
    max_col = ws3.max_column
    # Columns from E (5) to the end
    for c in range(5, max_col + 1):
        training_no = _to_str(ws3.cell(row=6, column=c).value)             # header row with Training_No.
        marker      = _to_str(ws3.cell(row=row_index, column=c).value)     # group's row marker (e.g. '1' / '0')

        if training_no is None:
            continue                  # no training header
        if marker is None:
            continue                  # treat blank as '0'
        if marker != "0":             # any non-zero / non-blank means "required"
            required.append(training_no)

    return required

# Iterate OVERVIEW (B=Employee_ID, E=Training_Group, I=Exit_Date, J=In_Maternity_Leave)
xlUp = -4162  # Excel constant
row_o = 3
added_total = 0

while True:
    emp_id_val = overview.Cells(row_o, 2).Value  # column B
    if emp_id_val in (None, ""):
        break

    emp_id_str = _to_str(emp_id_val) or ''
    exit_val   = overview.Cells(row_o, 9).Value  # column I: Exit_Date
    mat_leave  = overview.Cells(row_o, 10).Value # column J: In_Maternity_Leave
    group_val  = _to_str(overview.Cells(row_o, 5).Value)  # column E: Training_Group

    # Skip if Exit_Date present OR In_Maternity_Leave == YES  (Response 1 behavior)
    if (exit_val not in (None, "")) or (str(mat_leave).strip().upper() == "YES"):
        log(f"Skipping {emp_id_str}: Exit='{exit_val}' Maternity='{mat_leave}'")
        row_o += 1
        continue

    required_list = get_required_training_numbers(group_val)
    if not required_list:
        log(f"  No required trainings resolved for {emp_id_str} (group={group_val}); skipping.")
        row_o += 1
        continue

    # Get the employee's sheet (named with Employee_ID)
    try:
        emp_sheet = wb.Sheets(emp_id_str)
    except Exception as e:
        log(f"  WARNING: No sheet for {emp_id_str}: {e}")
        row_o += 1
        continue

    # Collect existing Training_No. from column A (A8 downward)
    last_row = emp_sheet.Cells(emp_sheet.Rows.Count, 1).End(xlUp).Row
    if last_row < 8:
        last_row = 7

    existing = set()
    for r in range(8, last_row + 1):
        v = emp_sheet.Cells(r, 1).Value  # column A
        s = _to_str(v)
        if s:
            existing.add(s)

    # Determine missing required training numbers
    missing = [t for t in required_list if t not in existing]

    if missing:
        log(f"  {emp_id_str}: adding {len(missing)} missing Training_No.: {missing}")    
    
     
    # Append each missing Training_No. as a new row with I="NOK", but only if training is ACTIVE
    added_this_emp = 0
    inactive_suppressed = 0
    # skipped_list = []  # <-- optional: collect training numbers skipped due to inactive flag
    
    for t in missing:
        # Only "1" is active; anything else (0/2/blank) => inactive → don't add NOK
        if ACTIVE_FLAG_BY_TRAINING.get(_to_str(t), "1") != "1":
            inactive_suppressed += 1
            # skipped_list.append(_to_str(t))  # <-- optional
            continue
    
        last_row += 1
        emp_sheet.Cells(last_row, 1).Value = t       # Column A: Training No.
        emp_sheet.Cells(last_row, 9).Value = "NOK"   # Column I: Status
        added_this_emp += 1
    
    # log line if anything was skipped by inactive flag
    if inactive_suppressed:
        log(f"  {emp_id_str}: {inactive_suppressed} missing item(s) skipped due to inactive training flag")
        # log(f"    skipped Training_No.: {', '.join(skipped_list)}")  # <-- optional detail
    
    added_total += added_this_emp
    row_o += 1
        
    log(f"Completed matrix backfill. Total new 'NOK' rows added: {added_total}")


# *** STEP 4.9: Add 'log' sheet with run data, status, and LOGfile ***
try:
    import os, getpass, socket
    from datetime import datetime

    # Identify script/user/IP
    script_name = os.path.basename(__file__) if "__file__" in globals() else "HR_SCRIPT but could not get exact name"
    user_name = getpass.getuser()
    ip_addr = socket.gethostbyname(socket.gethostname())

    # Compute end time for the sheet and runtime (does not change your STEP 6 logs)
    end_ts_log = datetime.now()
    duration_log = end_ts_log - start_time

    # Determine Status + Error message
    # If you set RUN_STATUS anywhere like: RUN_STATUS = {"status": "Failure", "error": str(e)}
    # this will pick it up. Otherwise we infer from log_lines (look for "ERROR"/"Traceback").
    rs = globals().get("RUN_STATUS", {})
    status = (rs.get("status") or "").strip() if isinstance(rs, dict) else ""
    error_msg = (rs.get("error") or "").strip() if isinstance(rs, dict) else ""
    if not status:
        status = "Success"
    if not error_msg:
        for _line in reversed(log_lines):
            if not _line:
                continue
            if "ERROR" in _line or "Traceback" in _line:
                status = "Failure"
                error_msg = _line
                break

    # Recreate 'log' sheet right after OVERVIEW
    try:
        wb.Sheets("log").Delete()
    except Exception:
        pass
    overview_ws = wb.Sheets("OVERVIEW")
    log_ws = wb.Sheets.Add(After=overview_ws)
    log_ws.Name = "log"

    # Key-value info
    info_rows = [
        ("Script name", script_name),
        ("Start time", start_time.strftime("%Y-%m-%d %H:%M:%S") if hasattr(start_time, "strftime") else str(start_time)),
        ("End time",   end_ts_log.strftime("%Y-%m-%d %H:%M:%S")),
        ("Total runtime", str(duration_log)),
        ("User", user_name),
        ("IP-Adress", ip_addr),  # keep your spelling
        ("Status", status),
        ("Error", error_msg),
        ("Others", ""),
    ]

    r = 1
    for k, v in info_rows:
        log_ws.Cells(r, 1).Value = k
        log_ws.Cells(r, 2).Value = "" if v is None else str(v)
        r += 1

    # LOGfile section
    log_ws.Cells(r + 1, 1).Value = "LOGfile (line by line)"
    rr = r + 2
    for line in log_lines:
        log_ws.Cells(rr, 1).Value = line
        rr += 1

    # Simple formatting
    log_ws.Columns("A:B").AutoFit()

except Exception as _e:
    # Do not fail the whole run if the 'log' sheet creation has issues
    log(f"WARNING: could not create 'log' sheet: {_e}")



# *** STEP 5: Save the new workbook with timestamp ***
log("*** STEP 5: Saving and closing...")
log(f"Saving final workbook to: {output_path}")
wb.SaveAs(output_path)
wb.Close(SaveChanges=True)
excel.Quit()

# Close COM objects --> if the script stop it's necessary to go into the Windows Task Manager, search for "Excel" and click: END TASK
del wb
del excel

# *** STEP 6: Finalize and write log ***
end_time = datetime.now()
duration = end_time - start_time
log("All sheets created and Excel closed.")
log(f"End Time: {end_time.strftime('%Y-%m-%d %H:%M:%S')}")
log(f"Total Duration: {duration}")

# *** STEP 7: Write log file ***
with open(log_path, "w", encoding="utf-8") as f:
    for line in log_lines:
        f.write(line + "\n")
print("Log written to:", log_path)