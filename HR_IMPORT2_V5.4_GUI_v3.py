#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
HR_IMPORT4 Rev 5.4  (GUI Removed)
Date: 2025-08-26 06:08:02
"""

import os
import sys
import logging
import re
from typing import List, Optional, Tuple, Dict
from datetime import datetime, date, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Font
from openpyxl.styles import Alignment, numbers
from openpyxl.utils import get_column_letter

# =========================
# CONFIGURATION
# =========================
USER_HOME = os.path.expanduser("~")

# Automatically detect which RRC SharePoint folder structure is synced
if os.path.exists(USER_HOME + r"\RRC power solutions\RRC VN - 999_SHARE_VN"):
    BASE_DIR = USER_HOME + r"\RRC power solutions\RRC VN - 999_SHARE_VN"
    print("User synced the 999_SHARE_VN folder separately!")
elif os.path.exists(USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN"):
    BASE_DIR = USER_HOME + r"\RRC power solutions\RRC VN - Documents\999_SHARE_VN"
    print("User synced the whole Sharepoint library Documents folder!")
else:
    raise FileNotFoundError("Cannot find the RRC 999_SHARE_VN folder.")

FILE1 = BASE_DIR + r"\280_HR\004_Training\100_Training\Training_Evidence_Scan.xlsx"
FILE4 = BASE_DIR + r"\280_HR\004_Training\100_Training\Training_Evidence_Scan\XLSX_to_import\Training_Evidence.xlsx"
FILE2 = BASE_DIR + r"\280_HR\004_Training\100_Training\Training_Matrix.xlsx"
FILE3 = r"C:\Training HR scripts\RRC_Employee_Master_Data_2025-07-10.xlsx"

# Logging to same folder as File1
LOG_FILE = os.path.join(os.path.dirname(FILE1), "HR_IMPORT4_Rev5_4.log")

COL_PRIMARY_KEY = 1     # A
COL_TRAIN_EVID_NO = 2   # B
COL_FILENAME = 3        # C (for hyperlink building in D)
COL_HYPERLINK = 4       # D
COL_TRAIN_NO = 5        # E
COL_TRAIN_START = 6     # F ("Training_date")
COL_VECTOR_VALUE = 7    # G
COL_EXPIRY = 8          # H
COL_STATUS = 9          # I
COL_TRAIN_CONTENT = 11  # K ('Training_Content')
COL_TRAIN_GROUP_NO = 13 # M ('Training Group No.')
COL_ACTIVE_FLAG = 15    # O
COL_FOLLOWUP = 16       # P
COL_IMPORT_DATE = 17    # Q

# File4 (source) cells/areas
SRC_TRAIN_NO_CELL = "C1"
SRC_TRAIN_START_CELL = "C9"
SRC_TRAIN_CONTENT_CELL = "C3"
SRC_EMPLOYEE_ID_COL = "B"    # start at B21 downward
SRC_EMPLOYEE_ID_START_ROW = 21   # start at B21 downward

# =========================
# LOGGING
# =========================
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
)
console = logging.StreamHandler()
console.setLevel(logging.INFO)
console.setFormatter(logging.Formatter("%(message)s"))
logging.getLogger().addHandler(console)

# =========================
# UTILITIES
# =========================
def parse_datetime(value) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y  %H:%M:%S", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None

def parse_training_no_list(value) -> List[int]:
    if value is None or str(value).strip() == "":
        return []
    parts = [p.strip() for p in str(value).split(",")]
    nums = []
    for p in parts:
        if p:
            try:
                nums.append(int(float(p)))
            except ValueError:
                pass
    return nums

def get_last_numeric_in_column(ws: Worksheet, col_idx: int) -> int:
    max_val = 0
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
        v = row[0]
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            try:
                iv = int(v)
                if iv > max_val:
                    max_val = iv
            except Exception:
                pass
    return max_val

def read_employee_ids(ws: Worksheet) -> List[str]:
    ids = []
    row = SRC_EMPLOYEE_ID_START_ROW
    while True:
        cell = ws[f"{SRC_EMPLOYEE_ID_COL}{row}"]
        val = cell.value
        if val is None or str(val).strip() == "":
            break
        ids.append(str(val).strip())
        row += 1
    return ids

def canonical_key_tuple3(training_no, training_dt, emp_id) -> Tuple[str, str, str]:
    def clean_numlike(x):
        if x is None:
            return ""
        s = str(x).strip()
        if s.endswith(".0"):
            s = s[:-2]
        return s

    if isinstance(training_dt, datetime):
        dt_key = training_dt.replace(microsecond=0).strftime("%Y-%m-%d %H:%M:%S")
    else:
        dt = parse_datetime(training_dt)
        if dt is None:
            dt_key = str(training_dt).strip()
        else:
            dt_key = dt.replace(microsecond=0).strftime("%Y-%m-%d %H:%M:%S")
            
    return (clean_numlike(training_no), dt_key, clean_numlike(emp_id))

def clean_decimal_str(x) -> str:
    if x is None or str(x).strip() == "":
        return ""
    s = str(x).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s

def extract_evidence_no_from_sheet_name(ws_name: str):
    name = os.path.basename(str(ws_name)).strip()
    m = re.match(r'(?i)^Training_(.+?)\.pdf$', name)
    if not m:
        raise ValueError(f"Cannot derive Training Evidence No. from sheet name: {ws_name!r}")
    evid = m.group(1).strip()
    try:
        return int(evid)
    except ValueError:
        return evid

# =========================
# STEP 0: Load File1 and prepare structures
# =========================
def step0_load_file1():
    logging.info("=== START (Rev 5.4) ===")
    logging.info(f"File1 (base): {FILE1}")
    logging.info(f"File4 (source): {FILE4}")
    logging.info(f"File2 (matrix): {FILE2}")
    logging.info(f"File3 (employee master): {FILE3}")
    logging.info(f"Log file: {LOG_FILE}")
    if not os.path.exists(FILE1):
        logging.error(f"File1 not found: {FILE1}")
        sys.exit(1)
    if not os.path.exists(FILE4):
        logging.error(f"File4 not found: {FILE4}")
        sys.exit(1)
    if not os.path.exists(FILE2):
        logging.error(f"File2 not found: {FILE2}")
        sys.exit(1)
    if not os.path.exists(FILE3):
        logging.error(f"File3 not found: {FILE3}")
        sys.exit(1)

    wb1 = load_workbook(FILE1)
    ws1 = wb1.active  
    return wb1, ws1

# =========================
# STEP 1: Build existing keys
# =========================
def step1_build_existing_keyset(ws1: Worksheet) -> Tuple[set, int, int]:
    keyset = set()
    last_pk = get_last_numeric_in_column(ws1, COL_PRIMARY_KEY)
    last_evid = get_last_numeric_in_column(ws1, COL_TRAIN_EVID_NO)

    max_row = ws1.max_row
    for r in range(2, max_row + 1):
        tno = ws1.cell(row=r, column=COL_TRAIN_NO).value
        tdt = ws1.cell(row=r, column=COL_TRAIN_START).value
        eid = ws1.cell(row=r, column=12).value  # L: Employee_ID
        
        if tno is None and eid is None:
            continue
            
        key = canonical_key_tuple3(tno, tdt, eid)
        keyset.add(key)

    logging.info(f"Detected last Primary Key in File1: {last_pk}")
    logging.info(f"Detected last Training Evidence No. in File1: {last_evid}")
    return keyset, last_pk, last_evid

# =========================
# STEP 2: Dry-run & ROW-LEVEL Dupe Check
# =========================
def step2_dry_run_and_dupe_check(ws1: Worksheet, last_pk: int, last_evid: int) -> Tuple[list, int, int]:
    from openpyxl import load_workbook as _lw
    wb_src = _lw(FILE4, data_only=True)
    candidates = []
    pk = last_pk
    evid = last_evid

    existing_keys, _, _ = step1_build_existing_keyset(ws1)
    import_ts = datetime.now()

    skipped_sheets_count = 0
    imported_sheets_count = 0

    for ws_name in wb_src.sheetnames:
        ws_src = wb_src[ws_name]
        
        if ws_name.upper() in ["SUMMARY", "TEMPLATE", "SHEET1"]:
            continue
            
        logging.info(f"-- Planning import for source sheet: {ws_name} --")

        training_nos = parse_training_no_list(ws_src[SRC_TRAIN_NO_CELL].value)
        training_dt  = parse_datetime(ws_src[SRC_TRAIN_START_CELL].value)
        training_content = ws_src[SRC_TRAIN_CONTENT_CELL].value
        employee_ids = read_employee_ids(ws_src)

        if not training_nos:
            logging.info(f"[{ws_name}] No Training_No. in {SRC_TRAIN_NO_CELL}. Skip.")
            continue
        if training_dt is None:
            logging.info(f"[{ws_name}] Training_date in {SRC_TRAIN_START_CELL} invalid. Skip.")
            continue
        if not employee_ids:
            logging.info(f"[{ws_name}] No Employee_IDs from {SRC_EMPLOYEE_ID_COL}{SRC_EMPLOYEE_ID_START_ROW}. Skip.")
            continue

        sheet_rows = []
        for tno in training_nos:
            for emp in employee_ids:
                key = canonical_key_tuple3(tno, training_dt, emp)
                if key in existing_keys:
                    logging.warning(f"  -> DUPLICATE ROW SKIPPED: Training_No={key[0]}, Date={key[1]}, Employee_ID={key[2]} already in File1.")
                    continue
                sheet_rows.append((int(tno), str(emp).strip(), key))

        if not sheet_rows:
            logging.warning(f"[{ws_name}] All trainees in this sheet were duplicates. Skipping entire sheet.")
            skipped_sheets_count += 1
            continue

        per_sheet_evid = extract_evidence_no_from_sheet_name(ws_name)
        logging.info(f"[{ws_name}] Derived Training Evidence No. from sheet name: {per_sheet_evid}")

        for tno, emp, key in sheet_rows:
            pk += 1
            candidates.append({
                "pk": pk,
                "evidence_no": per_sheet_evid,
                "evidence_file": str(per_sheet_evid) + ".pdf",
                "training_no": int(tno),
                "training_dt": training_dt,
                "emp_id": str(emp).strip(),
                "training_content": training_content,
                "import_ts": import_ts
            })
            existing_keys.add(key)

        imported_sheets_count += 1

    logging.info(f"Planned new rows to insert: {len(candidates)}")
    logging.info(f"Source sheets imported/partially imported: {imported_sheets_count}")
    logging.info(f"Source sheets completely skipped due to 100% duplicates: {skipped_sheets_count}")

    return candidates, pk, evid

# =========================
# STEP 3: Append rows
# =========================
def step3_append_rows(ws1: Worksheet, candidates: list):
    if ws1.cell(row=1, column=COL_IMPORT_DATE).value is None:
        ws1.cell(row=1, column=COL_IMPORT_DATE, value="Import date")
    if ws1.cell(row=1, column=COL_TRAIN_CONTENT).value is None:
        ws1.cell(row=1, column=COL_TRAIN_CONTENT, value="Training_Content")

    first_new_row = ws1.max_row + 1
    for item in candidates:
        next_row = ws1.max_row + 1
        if next_row < 2:
            next_row = 2

        ws1.cell(row=next_row, column=COL_PRIMARY_KEY, value=item["pk"])                
        ws1.cell(row=next_row, column=COL_TRAIN_EVID_NO, value=item["evidence_no"])     
        ws1.cell(row=next_row, column=COL_FILENAME, value=item["evidence_file"])        
        ws1.cell(row=next_row, column=COL_TRAIN_NO, value=item["training_no"])          
        cF = ws1.cell(row=next_row, column=COL_TRAIN_START, value=item["training_dt"])  
        cF.number_format = "dd.mm.yyyy hh:mm:ss"

        ws1.cell(row=next_row, column=12, value=item["emp_id"])                         
        ws1.cell(row=next_row, column=COL_TRAIN_CONTENT, value=item["training_content"])  
        cQ = ws1.cell(row=next_row, column=COL_IMPORT_DATE, value=item["import_ts"])    
        cQ.number_format = "dd.mm.yyyy  hh:mm:ss"  

    last_new_row = ws1.max_row
    logging.info(f"Appended {len(candidates)} rows to File1. New rows: {first_new_row}-{last_new_row}.")
    return first_new_row, last_new_row

# =========================
# STEP 4: Support Data
# =========================
def step4_load_support_data() -> Tuple[Dict[str, str], pd.DataFrame, Dict[str, object]]:
    df3 = pd.read_excel(FILE3, sheet_name="MASTER_DATA", header=None, skiprows=4, dtype={0: str})
    df3 = df3[[0, 4]]
    df3.columns = ["Employee_ID", "Training_Group_No"]
    df3["Employee_ID"] = df3["Employee_ID"].astype(str).str.strip()
    
    def _clean(x):
        try:
            if pd.isnull(x) or x == "":
                return ""
            s = str(x).strip()
            if s.endswith(".0"):
                s = s[:-2]
            return s
        except Exception:
            return str(x)
            
    df3["Employee_ID"] = df3["Employee_ID"].apply(_clean)
    df3["Training_Group_No"] = df3["Training_Group_No"].apply(_clean)
    empid_to_group = dict(zip(df3["Employee_ID"], df3["Training_Group_No"]))

    matrix_df = pd.read_excel(FILE2, sheet_name="MATRIX", header=5, usecols="D:FF")
    matrix_df.columns = matrix_df.columns.map(str)
    matrix_df = matrix_df.dropna(how="all")

    def normalize_training_no(value) -> str:
        if value is None:
            return ""
        s = str(value).strip().replace("\u00a0", "")
        if not s:
            return ""
        s = s.replace(",", ".")
        if s.endswith(".0"):
            s = s[:-2]
        return s
    
    wb2 = load_workbook(FILE2, data_only=True, read_only=True)
    ws2 = wb2["MATRIX"]

    active_flag_map: Dict[str, object] = {}
    for col in range(5, ws2.max_column + 1):
        training_no_key = normalize_training_no(ws2.cell(row=6, column=col).value)
        if not training_no_key:
            continue
        if training_no_key not in active_flag_map:
            active_flag_map[training_no_key] = ws2.cell(row=4, column=col).value

    wb2.close()
    return empid_to_group, matrix_df, active_flag_map

# =========================
# STEP 5: Enrich ONLY new rows
# =========================
def step5_enrich_new_rows(ws1: Worksheet, first_new_row: int, last_new_row: int,
                          empid_to_group: Dict[str, str], matrix_df: pd.DataFrame,
                          active_flag_map: Dict[str, object]):
    def _to_clean_str(x):
        if x is None or str(x).strip() == "":
            return ""
        s = str(x).strip()
        if s.endswith(".0"):
            s = s[:-2]
        return s

    matrix_group_col = matrix_df.columns[0]
    matrix_index = { _to_clean_str(r[matrix_group_col]) : r for _, r in matrix_df.iterrows() }

    def compute_status(active_flag_val, vec_val, tdate_cell):
        v = tdate_cell
        if isinstance(v, datetime):
            tdate = v.date()
        elif isinstance(v, date):
            tdate = v
        else:
            return "OK" 
        try:
            vec = int(float(str(vec_val)))
        except Exception:
            vec = None
        if vec == 2:
            expiry = tdate + timedelta(days=365)
        elif vec == 3:
            expiry = tdate + timedelta(days=2*365)
        else:
            expiry = date(9999, 12, 31)
        try:
            af = int(float(str(active_flag_val)))
        except Exception:
            af = 0
        return "NOK" if (af == 1 and expiry < date.today()) else "OK"

    def normalize_key(v) -> str:
        if v is None:
            return ""
        s = str(v).strip().replace("\u00a0", "").replace(" ", "")
        if not s:
            return ""
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            s = s.replace(",", ".")
        if s.endswith(".0"):
            s = s[:-2]
        return s 

    def build_vector_lookup(file2_path: str):
        wb = load_workbook(file2_path, data_only=True, read_only=True)
        ws = wb["MATRIX"]
        row_lookup = {}
        col_lookup = {}

        for r in range(7, 63):
            key = normalize_key(ws.cell(r, 4).value)
            if key and key not in row_lookup:
                row_lookup[key] = r

        for c in range(5, ws.max_column + 1):  
            key = normalize_key(ws.cell(6, c).value)
            if key and key not in col_lookup: 
                col_lookup[key] = c

        return wb, ws, row_lookup, col_lookup

    if ws1.cell(row=1, column=COL_VECTOR_VALUE).value is None:
        ws1.cell(row=1, column=COL_VECTOR_VALUE, value="Vector_value from Training Matrix")
    if ws1.cell(row=1, column=COL_EXPIRY).value is None:
        ws1.cell(row=1, column=COL_EXPIRY, value="Expiration date calculated according Vector_value")
    if ws1.cell(row=1, column=COL_STATUS).value is None:
        ws1.cell(row=1, column=COL_STATUS, value="Training_Compliance_Status (calculated)")
    if ws1.cell(row=1, column=COL_ACTIVE_FLAG).value is None:
        ws1.cell(row=1, column=COL_ACTIVE_FLAG, value="Active Flag")
    if ws1.cell(row=1, column=COL_FOLLOWUP).value is None:
        ws1.cell(row=1, column=COL_FOLLOWUP, value="Follow up Training Evidence No.")
    if ws1.cell(row=1, column=COL_TRAIN_GROUP_NO).value is None:
        ws1.cell(row=1, column=COL_TRAIN_GROUP_NO, value="Training Group No.")
    if ws1.cell(row=1, column=COL_TRAIN_CONTENT).value is None:
        ws1.cell(row=1, column=COL_TRAIN_CONTENT, value="Training_Content")

    matrix_wb, matrix_ws, vector_row_map, vector_col_map = build_vector_lookup(FILE2)

    updated = 0
    for r in range(first_new_row, last_new_row + 1):
        emp_id = _to_clean_str(ws1.cell(row=r, column=12).value)
        tno    = _to_clean_str(ws1.cell(row=r, column=COL_TRAIN_NO).value)
        tdate  = ws1.cell(row=r, column=COL_TRAIN_START).value

        vec_value = ""
        group_no = empid_to_group.get(emp_id, "")
        row_idx = vector_row_map.get(normalize_key(group_no))
        col_idx = vector_col_map.get(tno)

        if row_idx is not None and col_idx is not None:
            vec_value = matrix_ws.cell(row=row_idx, column=col_idx).value

        ws1.cell(row=r, column=COL_VECTOR_VALUE, value=vec_value)
        ws1.cell(row=r, column=COL_TRAIN_GROUP_NO, value=group_no)

        fname = ws1.cell(row=r, column=COL_FILENAME).value
        if fname:
            formula = f'=HYPERLINK("https://rrcpowersolutions.sharepoint.com/sites/RRCVN/Shared Documents/999_SHARE_VN/280_HR/004_Training/100_Training/Training_Evidence_Scan/PDF_to_import/" & C{r}, "Open " & C{r})'
            ws1.cell(row=r, column=COL_HYPERLINK, value=formula)
            ws1.cell(row=r, column=COL_HYPERLINK).font = Font(color="0000FF", underline="single")

        ws1.cell(row=r, column=COL_EXPIRY).value = (
            f'=IF(OR(G{r}=0,G{r}=1),2958465,IF(G{r}=2,F{r}+365,IF(G{r}=3,F{r}+2*365,"")))'
        )
        ws1.cell(row=r, column=COL_EXPIRY).number_format = numbers.FORMAT_DATE_XLSX14

        active_flag = active_flag_map.get(tno, "")
        ws1.cell(row=r, column=COL_ACTIVE_FLAG, value=active_flag)

        status_val = compute_status(active_flag, vec_value, tdate)
        ws1.cell(row=r, column=COL_STATUS, value=status_val)

        updated += 1

    logging.info(f"Enriched {updated} newly imported rows (K,G,D,H,O,I,M).")

# =========================
# STEP 6: Link follow-ups 
# =========================
def step6_link_followups(ws1: Worksheet, first_new_row: int, last_new_row: int):
    latest_ok_by_key: Dict[Tuple[str, str], Tuple[int, date, int]] = {}
    for r in range(first_new_row, last_new_row + 1):
        emp_id = clean_decimal_str(ws1.cell(row=r, column=12).value)
        tno    = clean_decimal_str(ws1.cell(row=r, column=COL_TRAIN_NO).value)
        status = str(ws1.cell(row=r, column=COL_STATUS).value or "").strip().upper()
        evno   = ws1.cell(row=r, column=COL_TRAIN_EVID_NO).value
        tdatev = ws1.cell(row=r, column=COL_TRAIN_START).value
        
        if isinstance(tdatev, datetime):
            tdate_key = tdatev.date()
        elif isinstance(tdatev, date):
            tdate_key = tdatev
        else:
            tdate_key = date.min

        if status == "OK" and emp_id and tno:
            key = (emp_id, tno)
            prev = latest_ok_by_key.get(key)
            if prev is None or (tdate_key, r) > (prev[1], prev[2]):
                latest_ok_by_key[key] = (evno, tdate_key, r)

    if not latest_ok_by_key:
        logging.info("No new OK rows to link as follow-ups. Skipping STEP 6.")
        return

    updates = 0
    misses  = 0
    checked = 0
    for r in range(2, first_new_row):
        emp_id = clean_decimal_str(ws1.cell(row=r, column=12).value)
        tno    = clean_decimal_str(ws1.cell(row=r, column=COL_TRAIN_NO).value)
        status = str(ws1.cell(row=r, column=COL_STATUS).value or "").strip().upper()

        if status == "NOK":
            checked += 1
            ok_info = latest_ok_by_key.get((emp_id, tno))
            if ok_info is None:
                misses += 1
                continue
            evno_ok, _, _ = ok_info
            
            if not ws1.cell(row=r, column=COL_FOLLOWUP).value:
                ws1.cell(row=r, column=COL_FOLLOWUP, value=evno_ok)
            ws1.cell(row=r, column=COL_STATUS, value="Recertified")
            updates += 1

    logging.info(f"STEP 6 follow-up linking: checked {checked} old NOK rows, updated {updates}, no-match {misses}.")

# =========================
# STEP 7: Light formatting
# =========================
def step7_format(ws1: Worksheet):
    max_row = ws1.max_row
    max_col = ws1.max_column
    last_col_letter = get_column_letter(max_col)

    for c in range(1, max_col + 1):
        cell = ws1.cell(row=1, column=c)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    ws1.auto_filter.ref = f"A1:{last_col_letter}{max_row}"
    ws1.freeze_panes = "A2"

    for r in range(2, max_row + 1):
        if ws1[f"F{r}"].value is not None:
            ws1[f"F{r}"].number_format = numbers.FORMAT_DATE_DATETIME
        if ws1[f"H{r}"].value is not None:
            ws1[f"H{r}"].number_format = numbers.FORMAT_DATE_XLSX14

    for col_letter in list("ABCDEFGHIJ"):
        ws1.column_dimensions[col_letter].width = 14
    ws1.column_dimensions["D"].width = 18
    ws1.column_dimensions["K"].width = 50

# =========================
# MAIN
# =========================
def main():
    wb1, ws1 = step0_load_file1()
    existing_keys, last_pk, last_evid = step1_build_existing_keyset(ws1)

    candidates, new_last_pk, new_last_evid = step2_dry_run_and_dupe_check(ws1, last_pk, last_evid)
    if not candidates:
        logging.info("No candidates to import. Nothing to do.")
        return

    first_new_row, last_new_row = step3_append_rows(ws1, candidates)
    empid_to_group, matrix_df, active_flag_map = step4_load_support_data()
    
    step5_enrich_new_rows(ws1, first_new_row, last_new_row, empid_to_group, matrix_df, active_flag_map)
    step6_link_followups(ws1, first_new_row, last_new_row)
    step7_format(ws1)

    wb1.save(FILE1)
    logging.info(f"SAVED updates to File1: {FILE1}")
    logging.info("=== COMPLETED (Rev 5.4) ===")

if __name__ == "__main__":
    main()